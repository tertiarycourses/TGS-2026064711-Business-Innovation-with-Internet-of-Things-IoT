#!/usr/bin/env python3
"""Generate the Business Innovation with IoT (TGS-2026064711) Lesson Plan (LP) DOCX
in the Tertiary house format.

Cover page + Document Version Control Record + auto TOC + Arial 11pt body +
colour-coded schedule table + page numbers.

ONE training day of 8 hours (9:30am-6:30pm with a 1-hour lunch). The component
split is fixed by the approved Course Proposal (CA-WSQ-2020-007388):
    Classroom Facilitation   300 min  (5.0 h, tea breaks counted within)
    Practical / Practicum     90 min  (1.5 h — Topic 3 IoTFlow labs)
    Assessment                90 min  (1.5 h — WA 75 min + OQ 15 min, per the approved Assessment Plan)
    ----------------------------------
    Total training           480 min  (8.0 h, excluding the 1-hour lunch)

The build asserts this split, so the LP can never silently drift out of
alignment with the accredited proposal.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
def _find_repo(start):
    env=os.environ.get("COURSE_REPO")
    if env and os.path.isdir(env): return env
    d=start
    for _ in range(8):
        d=os.path.dirname(d)
        if os.path.isdir(os.path.join(d,"courseware")) and os.path.isdir(os.path.join(d,"labs")): return d
    return os.path.dirname(os.path.dirname(HERE))
REPO=_find_repo(HERE); ASSETS=os.path.join(os.path.dirname(HERE),"assets")

BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
HEADER_FILL="1F6FEB"; TOPIC_FILL="E8F0FE"; BREAK_FILL="FFF4E5"; LUNCH_FILL="FDE9D9"
ASSESS_FILL="E8F7EE"; LAB_FILL="E6FAF3"

LAB_LIST="; ".join(f"Lab {a['num']}: {a['title']}" for a in ACT)

# ------------------------------------------------ schedule (single source of truth for timing)
# (start, end, minutes, kind, activity_text)   kind: admin/topic/lab/break/lunch/assess
SCHEDULE = [
    ("9:30am","9:45am",15,"admin","Welcome, trainer and learner introductions, ground rules, and mandatory digital attendance (AM)"),
    ("9:45am","9:55am",10,"admin","Learning outcomes, course outline and Skills Framework (TSC ICT-SNA-4003-1.1) briefing"),
    ("9:55am","10:50am",55,"topic","Topic 1 — Overview of Internet of Things (IoT): Industry 4.0, IoT building blocks, sensors and actuators, wireless technologies, and the new business models IoT enables. Includes the facilitated discussion 'Compare Business Models'.  (LO1 / A1, K1)"),
    ("10:50am","11:05am",15,"break","Tea break"),
    ("11:05am","12:00pm",55,"topic","Topic 2 — Potential Opportunities of IoT Innovation: use cases across agriculture, autonomous vehicles, logistics, smart buildings, healthcare and drones, and a repeatable method for finding an opportunity. Includes the facilitated discussion 'Spot an IoT Opportunity'.  (LO2 / A2, K1)"),
    ("12:00pm","12:30pm",30,"topic","Topic 3 — IoT Technology (concepts, part 1): cloud computing for IoT, the IoTFlow platform, the IoT data round trip, and uplink versus downlink.  (LO3 / A3, K2)"),
    ("12:30pm","1:30pm",60,"lunch","Lunch break"),
    ("1:30pm","1:40pm",10,"admin","Mandatory digital attendance (PM)"),
    ("1:40pm","1:55pm",15,"topic","Topic 3 — IoT Technology (concepts, part 2): turning data into business insight — visualise, detect, notify, record, analyse with AI and act automatically."),
    ("1:55pm","3:25pm",90,"lab","Topic 3 — HANDS-ON PRACTICAL on IoTFlow (https://iot.tertiaryinfotech.com). "+LAB_LIST),
    ("3:25pm","3:40pm",15,"break","Tea break"),
    ("3:40pm","4:20pm",40,"topic","Topic 4 — IoT Adoption: adoption barriers, why IoT security differs, the Mirai botnet case study, edge versus cloud processing, and cost-benefit analysis. Includes the facilitated discussion 'Cost-Benefit Analysis'.  (LO4 / A4, K3)"),
    ("4:20pm","4:45pm",25,"topic","Topic 5 — IoT Implementation: the six success factors, the Stanford d.school design-thinking model, and the 7-step implementation action plan. Includes the facilitated discussion 'Your Implementation Action Plan'.  (LO5 / A5, K4)"),
    ("4:45pm","5:00pm",15,"admin","Course recap and Q&A, TRAQOM survey, Assessment digital attendance, and Briefing for Assessment"),
    ("5:00pm","6:15pm",75,"assess","Written Assessment (WA) — Short-Answer Questions (Q&A), 75 minutes, open book. Assesses A1–A5 (15 minutes per ability)."),
    ("6:15pm","6:30pm",15,"assess","Oral Questioning (OQ) — 15 minutes, 1:1 with the assessor, open book. Assesses K1–K4."),
]

# ------------------------------------------------ build document
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)

prodoc.add_cover_page(doc,"LESSON PLAN",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[
 ("9.0","1 January 2026","Previous release — 2-day structure with ThingSpeak/MATLAB and IFTTT practical activities.",C.TRAINER),
 ("10.0","8 August 2026",
  "Major revision. Restructured to the accredited one-day (8-hour) format per Course Proposal "
  "CA-WSQ-2020-007388. All practical activities replaced with four hands-on labs on the IoTFlow "
  "platform (https://iot.tertiaryinfotech.com); ThingSpeak, MATLAB and IFTTT content retired. "
  "Assessment stated as WA (Q&A) plus Oral Questioning.",C.TRAINER),
 (C.VERSION.lstrip("v"),C.VERSION_DATE,
  "Each lab is now a self-contained folder (labs/lab-NN/) holding its instructions plus runnable "
  "Python scripts built on the official iotflow client. No change to the schedule, timings or "
  "assessment.",C.TRAINER),
])
prodoc.add_toc(doc)

def H(text,level=1):
    return doc.add_heading(text,level=level)

H("Course Information",1)
info=[("Course Title",C.TITLE),("WSQ Course Reference",C.COURSE_CODE),
      ("Training Provider",C.ORG+"  ("+C.UEN.replace('UEN: ','UEN ')+")"),
      ("TSC Title / Code",f"{C.TSC_TITLE}  ·  {C.TSC_CODE}  (Level 4)"),
      ("Duration","1 day · 8 training hours (excluding a 1-hour lunch break)"),
      ("Daily Timing","9:30 am – 6:30 pm (1-hour lunch; tea breaks counted within training time)"),
      ("Course Components","Classroom Facilitation 5.0 h · Practical / Practicum 1.5 h · Assessment 1.5 h"),
      ("Mode","Instructor-led classroom facilitation with hands-on IoT labs on the IoTFlow cloud platform"),
      ("Platform",f"{C.PLATFORM} — {C.PLATFORM_URL}"),
      ("Trainer",C.TRAINER)]
t=doc.add_table(rows=0,cols=2); t.style="Table Grid"
for k,v in info:
    c=t.add_row().cells; c[0].text=""; r=c[0].paragraphs[0].add_run(k); r.bold=True; r.font.size=Pt(10)
    prodoc._shade_cell(c[0],TOPIC_FILL)
    c[1].text=""; c[1].paragraphs[0].add_run(v).font.size=Pt(10)

H("Learning Outcomes",1)
doc.add_paragraph("On completion of this course, learners will be able to:")
for lo in C.LEARNING_OUTCOMES:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(lo).font.size=Pt(10.5)

H("Skills Framework Mapping (TSC ICT-SNA-4003-1.1 Business Innovation)",1)
doc.add_paragraph(C.TSC_DESC)
tm=doc.add_table(rows=0,cols=2); tm.style="Table Grid"
hdr=tm.add_row().cells
for i,htext in enumerate(["Abilities (assessed by the Written Assessment)","Knowledge (assessed by Oral Questioning)"]):
    cell=hdr[i]; cell.text=""
    r=cell.paragraphs[0].add_run(htext); r.bold=True; r.font.size=Pt(10)
    r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF); prodoc._shade_cell(cell,HEADER_FILL)
for i in range(max(len(C.ABILITIES),len(C.KNOWLEDGE))):
    cells=tm.add_row().cells
    for j,src in enumerate((C.ABILITIES,C.KNOWLEDGE)):
        cells[j].text=""
        if i<len(src):
            code,txt=src[i]
            p=cells[j].paragraphs[0]
            rb=p.add_run(code+" — "); rb.bold=True; rb.font.size=Pt(9.5)
            p.add_run(txt).font.size=Pt(9.5)

H("Assessment",1)
for a in [C.ASSESSMENT["written"],C.ASSESSMENT["practical"],
          "Format: Open Book — course slides, Learner Guide and approved materials only.",
          "The assessment is conducted at the end of the training day, from 5:00 pm.",
          "Assessor-to-learner ratio: 1:3 to 1:10 for the WA (Q&A); 1:1 for Oral Questioning.",
          C.ASSESSMENT["note"]]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(a).font.size=Pt(10.5)

def set_cell(cell,text,bold=False,size=9.5,color=None,fill=None,align=None):
    cell.text=""; p=cell.paragraphs[0]
    if align: p.alignment=align
    r=p.add_run(text); r.bold=bold; r.font.size=Pt(size); r.font.name="Arial"
    if color: r.font.color.rgb=color
    if fill: prodoc._shade_cell(cell,fill)

KIND_FILL={"topic":TOPIC_FILL,"break":BREAK_FILL,"lunch":LUNCH_FILL,"assess":ASSESS_FILL,
           "admin":"F3F5F8","lab":LAB_FILL}

H("Course Schedule",1)
H(f"Day 1 — {C.DAY_THEMES[1]}",2)
tbl=doc.add_table(rows=0,cols=3); tbl.style="Table Grid"; tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
hdr=tbl.add_row().cells
for i,htext in enumerate(["Time","Duration","Topic / Activity"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
training=classroom=practical=assessment=0
for start,end,mins,kind,text in SCHEDULE:
    cells=tbl.add_row().cells; fill=KIND_FILL.get(kind)
    set_cell(cells[0],f"{start}–{end}",bold=(kind in ("topic","assess","lab")),size=9.5,fill=fill)
    set_cell(cells[1],f"{mins} min",size=9.5,fill=fill)
    set_cell(cells[2],text,bold=(kind in ("topic","assess","lab")),size=9.5,fill=fill)
    if kind!="lunch": training+=mins
    if kind in ("topic","admin","break"): classroom+=mins
    elif kind=="lab": practical+=mins
    elif kind=="assess": assessment+=mins
for row in tbl.rows:
    row.cells[0].width=Inches(1.15); row.cells[1].width=Inches(0.9); row.cells[2].width=Inches(4.75)
p=doc.add_paragraph()
r=p.add_run(f"Total training time: {training} minutes ({training//60} hours), excluding the 1-hour lunch break. "
            f"Classroom Facilitation {classroom} min ({classroom/60:.1f} h) · "
            f"Practical/Practicum {practical} min ({practical/60:.1f} h) · "
            f"Assessment {assessment} min ({assessment/60:.1f} h).")
r.italic=True; r.font.size=Pt(9.5); r.font.color.rgb=GREY

# Guard: the LP can never drift from the accredited component split.
assert training==480,   f"Training minutes = {training}, expected 480"
assert classroom==300,  f"Classroom minutes = {classroom}, expected 300"
assert practical==90,   f"Practical minutes = {practical}, expected 90"
assert assessment==90,  f"Assessment minutes = {assessment}, expected 90"

H("Topic and Activity Reference",1)
tt=doc.add_table(rows=0,cols=3); tt.style="Table Grid"
hdr=tt.add_row().cells
for i,htext in enumerate(["Topic / Learning Unit","Outcome & TSC mapping","In-class activity"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
DISC_BY_TOPIC={d["topic"]:d for d in C.DISCUSSIONS}
for tp in C.TOPICS:
    acts=[a for a in ACT if a["topic"]==tp["num"]]
    cells=tt.add_row().cells
    set_cell(cells[0],f"Topic {tp['code']}: {tp['title']}",bold=True,size=9.5,fill=TOPIC_FILL)
    set_cell(cells[1],tp["weighting"],size=9.5,fill=TOPIC_FILL)
    if acts:
        act=", ".join(f"Lab {a['num']}: {a['title']}" for a in acts)
    else:
        d=DISC_BY_TOPIC.get(tp["num"])
        act=f"{d['title']} ({d['mins']} min facilitated discussion)" if d else "Lecture and didactic questioning"
    set_cell(cells[2],act,size=9.5)

H("Hands-On Labs (Topic 3 — Practical / Practicum, 90 minutes)",1)
doc.add_paragraph(
    f"All labs run on {C.PLATFORM} ({C.PLATFORM_URL}). Full step-by-step instructions are in the "
    f"Learner Guide; the slide deck presents each lab's purpose and outcome only.")
tl=doc.add_table(rows=0,cols=3); tl.style="Table Grid"
hdr=tl.add_row().cells
for i,htext in enumerate(["Lab","Title","What the learner builds"]):
    set_cell(hdr[i],htext,bold=True,size=10,color=RGBColor(0xFF,0xFF,0xFF),fill=HEADER_FILL)
for a in ACT:
    cells=tl.add_row().cells
    set_cell(cells[0],f"Lab {a['num']}",bold=True,size=9.5,fill=LAB_FILL)
    set_cell(cells[1],a["title"],size=9.5)
    set_cell(cells[2],a["build"],size=9.5)
for row in tl.rows:
    row.cells[0].width=Inches(0.7); row.cells[1].width=Inches(2.5); row.cells[2].width=Inches(3.6)

H("Instructional Methods",1)
for m in ["Lecture — trainer-led explanation of concepts with visual slides.",
          "Didactic questioning — the trainer poses questions to check and deepen understanding.",
          "Facilitated discussion — learners apply each topic to their own organisation (Topics 1, 2, 4 and 5).",
          "Demonstration — the trainer demonstrates each IoTFlow workflow before the learners attempt it.",
          "Hands-on practical — learners complete Labs 1–4 on their own IoTFlow project (Topic 3)."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(m).font.size=Pt(10.5)

H("Resources Required",1)
for m in ["Classroom with projector or large screen, whiteboard and high-speed WiFi.",
          "One laptop per learner with a modern web browser (spare laptops are available).",
          f"A free learner account on {C.PLATFORM} ({C.PLATFORM_URL}).",
          "Python 3 with the IoTFlow client (pip install \"iotflow[mqtt]\") for the optional scripted labs.",
          "IoT hardware (ESP32/ESP8266) is optional — every lab can be completed from the browser and terminal.",
          "Course slides, Learner Guide and Lesson Plan available on https://lms-tms.tertiaryinfotech.com."]:
    p=doc.add_paragraph(style="List Bullet"); p.add_run(m).font.size=Pt(10.5)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
OUT=os.path.join(REPO,"courseware",f"LP-{C.SHORT_TITLE}.docx")
doc.save(OUT)
print("Saved",OUT)
print(f"  training={training} classroom={classroom} practical={practical} assessment={assessment}")
